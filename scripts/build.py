#!/usr/bin/env python3
"""
Auto OEM Trends Tracker — legacy build (workbook -> dashboard JSON).
===================================================================

Backward-compatible entry point: reads the two source workbooks in ``raw/`` and writes the
dashboard JSON into ``data/``. All of the parsing/finalization math now lives in
``lib/pipeline_core.py`` (shared with ``scripts/seed_history.py`` and ``lib/build.py``), so
this file is a thin orchestrator. Output is byte-for-byte identical to the original build.

    pip install openpyxl
    python3 scripts/build.py

Re-running is idempotent (it deterministically overwrites ``data/``). For the automated,
store-driven path used in production, see ``scripts/run.py`` -> ``lib/build.py``.
"""

from __future__ import annotations

import os
import sys
from collections import OrderedDict

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from lib import pipeline_core as pc  # noqa: E402

RAW = os.path.join(ROOT, "raw")
OUT = os.path.join(ROOT, "data")
SIAM_XLSX = os.path.join(RAW, "Monthly_SIAM_Industry_Data_Jun26.xlsx")
SPARK_XLSX = os.path.join(RAW, "Auto_Database_Summary__Spark.xlsx")


def parse_siam():
    wb = openpyxl.load_workbook(SIAM_XLSX, data_only=True, read_only=True)
    datasets = OrderedDict()
    for sheet, (cat_id, cat_label) in pc.SIAM_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ex = pc.extract_siam_sheet(wb[sheet])
        ds = pc.finalize_siam(cat_id, cat_label, ex["dates"], ex["names"], ex["raw"], ex["industry_raw"])
        datasets[f"{cat_id}__siam"] = ds
        print(f"  [SIAM] {sheet:20s} -> {cat_id:5s}  entities={len(ds['entities'])}  "
              f"months={len(ds['series']['monthly']['periods'])}  latest={ds['latest']['monthly']}")
    wb.close()
    return datasets


def parse_internal():
    wb = openpyxl.load_workbook(SPARK_XLSX, data_only=True, read_only=True)
    datasets = OrderedDict()

    dates, coll, _ev2w = pc.extract_spark_sheet1(wb["OEM - Summary - 2W, PV, 3W"])
    datasets.update(pc.build_internal_from_sheet1(dates, coll))

    t_dates, players, band_totals = pc.extract_spark_tractors(wb["Tractors"])
    datasets["tractors__internal"] = pc.finalize_tractors(t_dates, players, band_totals)

    quarters, seg_members, seg_totals = pc.extract_spark_cv(wb["CV"])
    datasets["cv__internal"] = pc.finalize_cv(quarters, seg_members, seg_totals)

    wb.close()
    for k, ds in datasets.items():
        base = ds["frequencies"][0]
        print(f"  [INTERNAL] {k:18s} entities={len(ds['entities'])}  "
              f"periods={len(ds['series'][base]['periods'])}  latest={ds['latest'].get(base)}  "
              f"ev={'yes' if ds.get('ev') else 'no'}  segs={len(ds.get('segments',{}).get(base,{}).get('groups',[]))}")
    return datasets


def main():
    print("Building Auto OEM Trends data ...")
    print("SIAM (industry):")
    siam = parse_siam()
    print("Internal DB (OEM-granular):")
    internal = parse_internal()

    datasets = OrderedDict()
    datasets.update(siam)
    datasets.update(internal)
    n, total_bytes = pc.emit(datasets, OUT)
    print(f"\n  Wrote manifest.json/.js + {n} dataset files (json+js) "
          f"({total_bytes/1024:.0f} KB data) to {OUT}")
    print("Done.")


if __name__ == "__main__":
    main()
