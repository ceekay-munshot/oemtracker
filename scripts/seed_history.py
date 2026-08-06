#!/usr/bin/env python3
"""
scripts/seed_history.py — one-time backfill of the two workbooks into the canonical store.
=========================================================================================

This preserves the full 1992->2026 history that the dashboard shipped with, by decomposing
each workbook into long-format records (see ``lib/store.py``):

  * ``raw/Monthly_SIAM_Industry_Data_Jun26.xlsx``  -> source ``SIAM``
  * ``raw/Auto_Database_Summary__Spark.xlsx``       -> source ``Internal-DB(historical)``

It reuses the exact parsing math from ``lib/pipeline_core`` (the same code the dashboard
build trusts), so what lands in the store re-derives byte-for-byte back to the committed
dashboard JSON via ``lib/build.py``.

Idempotent: re-running appends nothing (every record's signature already matches). Safe to
run repeatedly. Live months are added later by the adapters + ``scripts/run.py`` — never here.

    python3 scripts/seed_history.py
"""

from __future__ import annotations

import os
import sys
from datetime import datetime, timezone

import openpyxl

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)

from lib import model as M  # noqa: E402
from lib import pipeline_core as pc  # noqa: E402
from lib.store import Store, make_record  # noqa: E402

RAW = os.path.join(ROOT, "raw")
STORE_DIR = os.path.join(ROOT, "data", "store")
SIAM_XLSX = os.path.join(RAW, "Monthly_SIAM_Industry_Data_Jun26.xlsx")
SPARK_XLSX = os.path.join(RAW, "Auto_Database_Summary__Spark.xlsx")

SIAM_REF = "seed:raw/Monthly_SIAM_Industry_Data_Jun26.xlsx"
SPARK_REF = "seed:raw/Auto_Database_Summary__Spark.xlsx"


def _prov_set(periods, n_trailing):
    """The trailing N base periods are flagged provisional in the store (audit metadata)."""
    return set(periods[-n_trailing:]) if len(periods) >= n_trailing else set(periods)


def _emit(recs, store, ref, source, category, segment, oem, metric, freq, periods, values,
          prov_set, oem_raw=None, bracket=False):
    """
    Append one record per non-null value in ``values`` (aligned to ``periods``).

    ``bracket=True`` first applies ``clean_trailing_leading`` and stores only the resulting
    non-null (positive-bracketed) range. This is used for series the build cleans *directly*
    (every SIAM series, and internal *total* rows) — since cleaning is idempotent, the stored
    result re-derives identically, while all the meaningless leading/trailing zeros (e.g. an
    OEM's category it has never sold in) are dropped, keeping the store compact. It is NOT
    used for series that get summed with others before cleaning (internal sub-segment members),
    where a leading 0 (which contributes 0) must not be confused with a missing value.
    """
    if bracket:
        values, _ = pc.clean_trailing_leading(values)
    for i, v in enumerate(values):
        if v is None:
            continue
        recs.append(make_record(
            source, category, segment, oem, metric, freq, periods[i], v,
            oem_raw=oem_raw, provisional=(periods[i] in prov_set), source_ref=ref))


# --------------------------------------------------------------------------------------
# SIAM
# --------------------------------------------------------------------------------------

def seed_siam(store, ts):
    wb = openpyxl.load_workbook(SIAM_XLSX, data_only=True, read_only=True)
    total = 0
    for sheet, (cat, _label) in pc.SIAM_SHEETS.items():
        if sheet not in wb.sheetnames:
            continue
        ex = pc.extract_siam_sheet(wb[sheet])
        mkeys = [pc.month_key(d) for d in ex["dates"]]
        store.extend_axis(M.SRC_SIAM, cat, M.FREQ_M, mkeys)
        prov = _prov_set(mkeys, pc.PROVISIONAL_MONTHS)

        recs = []
        for eid, metricmap in ex["raw"].items():
            name = ex["names"][eid]
            for skey, vals in metricmap.items():
                metric = M.SIAM_KEY_TO_METRIC[skey]
                _emit(recs, store, SIAM_REF, M.SRC_SIAM, cat, M.SEG_ALL, name, metric,
                      M.FREQ_M, mkeys, vals, prov, bracket=True)
        for skey, vals in ex["industry_raw"].items():
            metric = M.SIAM_KEY_TO_METRIC[skey]
            _emit(recs, store, SIAM_REF, M.SRC_SIAM, cat, M.SEG_ALL, M.OEM_INDUSTRY, metric,
                  M.FREQ_M, mkeys, vals, prov, oem_raw="Industry", bracket=True)

        summ = store.upsert_many(recs, ingested_at=ts)
        total += summ["added"] + summ["restated"]
        print(f"  [SIAM] {sheet:20s} -> {cat:5s}  records+={summ['added']}  "
              f"restated={summ['restated']}  skipped={summ['skipped']}  months={len(mkeys)}")
    wb.close()
    return total


# --------------------------------------------------------------------------------------
# Internal DB (Spark)
# --------------------------------------------------------------------------------------

def seed_internal_sheet1(store, ts):
    wb = openpyxl.load_workbook(SPARK_XLSX, data_only=True, read_only=True)
    dates, coll, _ev = pc.extract_spark_sheet1(wb["OEM - Summary - 2W, PV, 3W"])
    wb.close()
    mkeys = [pc.month_key(d) for d in dates]
    prov = _prov_set(mkeys, pc.PROVISIONAL_MONTHS)
    src = M.SRC_INTERNAL_HIST

    # sub-segment members per category
    plan = {
        "2w": {"members": M.SEG_2W + [M.SEG_2W_EV], "seg_total": M.SEG_2W},
        "pv": {"members": M.SEG_PV, "seg_total": M.SEG_PV},
        "3w": {"members": M.SEG_3W, "seg_total": M.SEG_3W},
    }
    added = 0
    for cat, spec in plan.items():
        store.extend_axis(src, cat, M.FREQ_M, mkeys)
        recs = []
        # members
        for sub in spec["members"]:
            for name, vals in coll[cat][sub].items():
                _emit(recs, store, SPARK_REF, src, cat, sub, name, M.METRIC_DOMESTIC,
                      M.FREQ_M, mkeys, vals, prov)
        # segment totals
        for sub in spec["seg_total"]:
            vals = coll[cat]["seg_total"].get(sub)
            if vals is not None:
                _emit(recs, store, SPARK_REF, src, cat, sub, M.OEM_SEG_TOTAL, M.METRIC_DOMESTIC,
                      M.FREQ_M, mkeys, vals, prov, bracket=True)
        # EV total (2W only)
        if coll[cat].get("ev_total") is not None:
            _emit(recs, store, SPARK_REF, src, cat, M.SEG_2W_EV, M.OEM_EV_TOTAL, M.METRIC_DOMESTIC,
                  M.FREQ_M, mkeys, coll[cat]["ev_total"], prov, bracket=True)
        # explicit total-domestic (EV denominator)
        if coll[cat].get("dom_total") is not None:
            _emit(recs, store, SPARK_REF, src, cat, M.SEG_ALL, M.OEM_DOM_TOTAL, M.METRIC_DOMESTIC,
                  M.FREQ_M, mkeys, coll[cat]["dom_total"], prov, bracket=True)
        # explicit total-exports
        if coll[cat].get("exp_total") is not None:
            _emit(recs, store, SPARK_REF, src, cat, M.SEG_ALL, M.OEM_EXP_TOTAL, M.METRIC_EXPORTS,
                  M.FREQ_M, mkeys, coll[cat]["exp_total"], prov, bracket=True)

        summ = store.upsert_many(recs, ingested_at=ts)
        added += summ["added"] + summ["restated"]
        print(f"  [INTERNAL] sheet1 {cat:4s}  records+={summ['added']}  "
              f"restated={summ['restated']}  skipped={summ['skipped']}")
    return added


def seed_internal_tractors(store, ts):
    wb = openpyxl.load_workbook(SPARK_XLSX, data_only=True, read_only=True)
    dates, players, band_totals = pc.extract_spark_tractors(wb["Tractors"])
    wb.close()
    mkeys = [pc.month_key(d) for d in dates]
    prov = _prov_set(mkeys, pc.PROVISIONAL_MONTHS)
    src = M.SRC_INTERNAL_HIST
    store.extend_axis(src, "tractors", M.FREQ_M, mkeys)

    recs = []
    for name, vals in players.items():
        _emit(recs, store, SPARK_REF, src, "tractors", M.SEG_ALL, name, M.METRIC_DOMESTIC,
              M.FREQ_M, mkeys, vals, prov, bracket=True)
    for band, vals in band_totals.items():
        _emit(recs, store, SPARK_REF, src, "tractors", band, M.OEM_BAND_TOTAL, M.METRIC_DOMESTIC,
              M.FREQ_M, mkeys, vals, prov, bracket=True)

    summ = store.upsert_many(recs, ingested_at=ts)
    print(f"  [INTERNAL] tractors     records+={summ['added']}  "
          f"restated={summ['restated']}  skipped={summ['skipped']}")
    return summ["added"] + summ["restated"]


def seed_internal_cv(store, ts):
    wb = openpyxl.load_workbook(SPARK_XLSX, data_only=True, read_only=True)
    quarters, seg_members, seg_totals = pc.extract_spark_cv(wb["CV"])
    wb.close()
    q_keys = [pc.quarter_key(fy, q) for fy, q in quarters]
    prov = _prov_set(q_keys, 1)
    src = M.SRC_INTERNAL_HIST
    store.extend_axis(src, "cv", M.FREQ_Q, q_keys)

    recs = []
    for sid in pc.CV_SEGMENTS:
        for name, vals in seg_members[sid].items():
            _emit(recs, store, SPARK_REF, src, "cv", sid, name, M.METRIC_DOMESTIC,
                  M.FREQ_Q, q_keys, vals, prov)
        if seg_totals.get(sid) is not None:
            _emit(recs, store, SPARK_REF, src, "cv", sid, M.OEM_SEG_TOTAL, M.METRIC_DOMESTIC,
                  M.FREQ_Q, q_keys, seg_totals[sid], prov, bracket=True)

    summ = store.upsert_many(recs, ingested_at=ts)
    print(f"  [INTERNAL] cv           records+={summ['added']}  "
          f"restated={summ['restated']}  skipped={summ['skipped']}  quarters={len(q_keys)}")
    return summ["added"] + summ["restated"]


def main():
    if not (os.path.exists(SIAM_XLSX) and os.path.exists(SPARK_XLSX)):
        print("ERROR: source workbooks not found in raw/ — cannot seed.", file=sys.stderr)
        sys.exit(1)

    ts = datetime.now(timezone.utc).isoformat()
    store = Store(STORE_DIR)

    print("Seeding canonical store from workbooks ...")
    print(f"SIAM -> source '{M.SRC_SIAM}':")
    n1 = seed_siam(store, ts)
    print(f"Internal DB -> source '{M.SRC_INTERNAL_HIST}':")
    n2 = seed_internal_sheet1(store, ts)
    n3 = seed_internal_tractors(store, ts)
    n4 = seed_internal_cv(store, ts)

    store.save_axis()
    stats = store.write_stats()
    total_new = n1 + n2 + n3 + n4
    print(f"\n  Seed complete. New/restated records this run: {total_new}")
    for src, c in stats["sources"].items():
        print(f"    {src:26s} records={c['records']:>7d}  unique-keys={c['latest_keys']:>7d}")
    if total_new == 0:
        print("  (store already up to date — idempotent no-op)")


if __name__ == "__main__":
    main()
